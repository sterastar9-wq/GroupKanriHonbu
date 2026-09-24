#!/usr/bin/env python3
"""Build a traceable MF attendance candidate CSV from a Rakuco work-time XLSX.

This script intentionally creates a candidate only. A live Rakuco reviewer and
an independent CSV reviewer must pass before it can be released as final.csv.
"""

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
from pathlib import Path

import openpyxl


OUTPUT_HEADERS = ["従業員番号", "苗字", "名前", "打刻所属日", "打刻日", "打刻時間", "打刻種別"]
SOURCE_HEADERS = [
    "利用者名", "受給者証番号", "日付", "サービス提供状況", "通所開始時間", "通所終了時間",
    "食事提供", "送迎", "作業開始時間", "作業終了時間", "作業合計時間", "休憩合計時間",
    "在宅ワーク", "メモ", "作業内容",
]


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def as_text(value):
    return "" if value is None else str(value).strip()


def canonical_certificate(value):
    text = as_text(value)
    return text.zfill(10) if text.isdigit() and len(text) < 10 else text


def canonical_date(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%Y-%m-%d")
    match = re.search(r"(\d{2,4})/(\d{1,2})(?:/(\d{1,2}))?", as_text(value))
    if not match:
        raise ValueError(f"Invalid source date: {value!r}")
    year = match.group(1)
    if len(year) == 2:
        year = "20" + year
    if match.group(3) is None:
        raise ValueError(f"Source date has no day: {value!r}")
    return f"{year}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"


def canonical_time(value):
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    text = as_text(value)
    if text in ("", "-"):
        return None
    match = re.fullmatch(r"(\d{1,2}):(\d{2})", text)
    if not match:
        raise ValueError(f"Invalid source time: {value!r}")
    return f"{int(match.group(1)):02d}:{match.group(2)}"


def find_detail_table(workbook):
    for sheet in workbook.worksheets:
        title_row = None
        for row in sheet.iter_rows():
            if as_text(row[0].value) == "<利用者毎の詳細情報>":
                title_row = row[0].row
                break
        if not title_row:
            continue
        for row_number in range(title_row + 1, title_row + 8):
            headers = [as_text(sheet.cell(row_number, col).value) for col in range(1, 16)]
            if headers == SOURCE_HEADERS:
                return sheet, row_number
        raise ValueError(f"SOURCE_SCHEMA_MISMATCH: detail title found in {sheet.title}, but A:O headers differ")
    raise ValueError("SOURCE_SCHEMA_MISMATCH: <利用者毎の詳細情報> not found")


def read_master(path):
    result = {}
    with Path(path).open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"受給者証番号", "MF従業員番号", "姓", "名"}
        if not reader.fieldnames or not required.issubset(reader.fieldnames):
            raise ValueError("MASTER_SCHEMA_MISMATCH")
        for row in reader:
            certificate = canonical_certificate(row["受給者証番号"])
            if not certificate:
                continue
            if certificate in result:
                raise ValueError(f"AMBIGUOUS_MASTER_CERTIFICATE: {certificate}")
            result[certificate] = row
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx")
    parser.add_argument("master_snapshot_csv")
    parser.add_argument("output_dir")
    parser.add_argument("--target-month", required=True, help="YYYY-MM")
    parser.add_argument("--exclude-certificate", action="append", default=[])
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    master = read_master(args.master_snapshot_csv)
    excluded = {canonical_certificate(value) for value in args.exclude_certificate}
    workbook = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    sheet, header_row = find_detail_table(workbook)

    candidate, expected, exceptions = [], [], []
    starts = ends = 0
    for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not values[0]:
            continue
        raw = dict(zip(SOURCE_HEADERS, values[:15]))
        certificate = canonical_certificate(raw["受給者証番号"])
        date = canonical_date(raw["日付"])
        if not date.startswith(args.target_month + "-"):
            raise ValueError(f"OUT_OF_SCOPE_DATE row={row_number} date={date}")
        person = master.get(certificate)
        if certificate in excluded:
            exceptions.append([row_number, certificate, raw["利用者名"], "EXCLUDED_BY_RUN_CONFIG"])
            continue
        if person is None:
            exceptions.append([row_number, certificate, raw["利用者名"], "MASTER_NOT_FOUND"])
            continue
        if not as_text(person["MF従業員番号"]):
            exceptions.append([row_number, certificate, raw["利用者名"], "UNASSIGNED_EMPLOYEE"])
            continue
        date_out = date.replace("-", "/")
        for column, event_type in (("作業開始時間", "出勤"), ("作業終了時間", "退勤")):
            event_time = canonical_time(raw[column])
            if event_time is None:
                continue
            event = [as_text(person["MF従業員番号"]), as_text(person["姓"]), as_text(person["名"]), date_out, date_out, event_time, event_type]
            expected.append([row_number, certificate] + event)
            candidate.append(event)
            if event_type == "出勤":
                starts += 1
            else:
                ends += 1

    with (output_dir / "candidate.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(OUTPUT_HEADERS)
        writer.writerows(candidate)
    with (output_dir / "expected_events.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["source_row", "受給者証番号"] + OUTPUT_HEADERS)
        writer.writerows(expected)
    with (output_dir / "excluded_unassigned.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["source_row", "受給者証番号", "利用者名", "reason"])
        writer.writerows(exceptions)

    manifest = {
        "status": "PENDING_LIVE_RECONCILIATION",
        "target_month": args.target_month,
        "source": {"xlsx": str(Path(args.xlsx).resolve()), "sha256": sha256(args.xlsx), "sheet": sheet.title, "header_row": header_row},
        "master_snapshot": {"path": str(Path(args.master_snapshot_csv).resolve()), "sha256": sha256(args.master_snapshot_csv)},
        "counts": {"candidate_rows": len(candidate), "clock_in": starts, "clock_out": ends, "excluded_or_unassigned_source_rows": len(exceptions)},
    }
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False))


if __name__ == "__main__":
    main()
